"""Generate real DOCX regression fixtures. Requires python-docx 1.2+."""
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from io import BytesIO
import base64
import json
from lxml import etree
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'demos'
OUT.mkdir(exist_ok=True)
FZ = '方正小标宋_GBK'
SONG = '宋体'
HEI = '黑体'

def font(target, name=None, *, east=None, theme=None):
    rpr = target._element.get_or_add_rPr()
    old = rpr.find(qn('w:rFonts'))
    if old is not None:
        rpr.remove(old)
    el = OxmlElement('w:rFonts')
    if name:
        for slot in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            el.set(qn('w:' + slot), name)
    if east:
        el.set(qn('w:eastAsia'), east)
    if theme:
        for slot in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme']:
            el.set(qn('w:' + slot), theme + ('EastAsia' if slot == 'eastAsiaTheme' else 'Ascii'))
    rpr.insert(0, el)

def clear_fonts(style):
    rpr = style._element.get_or_add_rPr()
    el = rpr.find(qn('w:rFonts'))
    if el is not None:
        rpr.remove(el)

def make_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.3)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)
    normal = doc.styles['Normal']
    font(normal, SONG)
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(9)
    title = doc.styles['Title']
    font(title, FZ)
    title.font.size = Pt(21)
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True
    for name in ['Heading 1', 'Heading 2']:
        style = doc.styles[name]
        font(style, HEI)
        style.font.size = Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = False
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.core_properties.author = 'Word Markdown Demo'
    doc.core_properties.last_modified_by = 'Word Markdown Demo'
    doc.core_properties.created = datetime(2026, 10, 4, 0, 0, 0)
    doc.core_properties.modified = datetime(2026, 10, 4, 0, 0, 0)
    return doc

def para(doc, text, face=None, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    if face:
        font(r, face)
    if bold:
        r.bold = True
    return p

def save(doc, filename, theme_patch=None):
    buf = BytesIO()
    doc.save(buf)
    source = ZipFile(BytesIO(buf.getvalue()))
    with ZipFile(OUT / filename, 'w', ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == 'word/theme/theme1.xml' and theme_patch:
                root = etree.fromstring(data)
                ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
                for group, face in theme_patch.items():
                    el = root.find('.//a:' + group, ns)
                    for key in ['latin', 'ea', 'cs']:
                        el.find('a:' + key, ns).set('typeface', face)
                    for extra in el.findall('a:font', ns):
                        if extra.get('script') in ['Hans', 'Hant']:
                            extra.set('typeface', face)
                data = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
            info = ZipInfo(entry.filename, date_time=(2026, 10, 4, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, data)

# Fixture 1: explicit run fonts, first occurrence, partial Markdown-like text.
doc = make_doc()
para(doc, '这段前言位于标题之前，合并后仍需保留。', SONG)
para(doc, '关于推进文档数字化整理的通知', FZ, 'Title')
p = doc.add_paragraph(style='Heading 1')
font(p.add_run('一、'), HEI)
font(p.add_run('工作目标'), 'SimHei')
font(p.add_run('  '), SONG)
para(doc, '将各部门材料统一转换为 Markdown，便于检索与归档。', SONG)
para(doc, '二、实施安排', HEI, 'Heading 1')
para(doc, '请于十月十五日前完成首批材料整理。', SONG)
para(doc, '这只是加粗的正文，不是二级标题。', SONG, bold=True)
para(doc, '再次出现的小标宋保留为正文', FZ)
para(doc, '# 这行不是标题；*星号*与<script>alert(1)</script>均为原文。', SONG)
para(doc, '负责人\t张三\n复核人\t李四', SONG)
save(doc, '01_直接字体.docx')

# Fixture 2: paragraph and character style chains, direct override.
doc = make_doc()
base = doc.styles.add_style('SmallSongBase', WD_STYLE_TYPE.PARAGRAPH)
base.base_style = doc.styles['Normal']
font(base, 'FZXiaoBiaoSong-B05')
doc.styles['Title'].base_style = base
clear_fonts(doc.styles['Title'])
para(doc, '公共服务资料归档工作指引', style='Title')
base_hei = doc.styles.add_style('HeiBase', WD_STYLE_TYPE.PARAGRAPH)
base_hei.base_style = doc.styles['Normal']
font(base_hei, 'SimHei')
hei_child = doc.styles.add_style('HeiChild', WD_STYLE_TYPE.PARAGRAPH)
hei_child.base_style = base_hei
hei_child.font.size = Pt(14)
para(doc, '一、收集范围', style='HeiChild')
para(doc, '收集正式通知、会议纪要及业务说明。', SONG)
charbase = doc.styles.add_style('HeiCharacterBase', WD_STYLE_TYPE.CHARACTER)
font(charbase, HEI)
char_child = doc.styles.add_style('HeiCharacterChild', WD_STYLE_TYPE.CHARACTER)
char_child.base_style = charbase
p = doc.add_paragraph()
p.add_run('二、归档要求').style = char_child
para(doc, '保留原始文件，按部门与日期命名。', SONG)
para(doc, '本段直接指定宋体，覆盖黑体样式，应保留为正文。', SONG, 'HeiChild')
p = doc.add_paragraph()
font(p.add_run('黑体部分'), HEI)
font(p.add_run('与宋体部分混排，不应成为标题。'), SONG)
save(doc, '02_样式继承.docx')

# Fixture 3: theme font references, inherited doc defaults and field result.
doc = make_doc()
font(doc.styles['Title'], theme='major')
font(doc.styles['Heading 1'], theme='minor')
para(doc, '主题字体识别测试说明', style='Title')
para(doc, '一、主题标题', style='Heading 1')
para(doc, '本篇标题通过 Word 主题字体定义，文字本身没有直接字体设置。', SONG)
para(doc, '二、可见内容', style='Heading 1')
p = doc.add_paragraph()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'HYPERLINK "https://example.com"')
r = OxmlElement('w:r')
t = OxmlElement('w:t')
t.text = '链接显示文字保留，不会联网访问目标。'
r.append(t)
field.append(r)
p._p.append(field)
para(doc, '手动编号 1. 保留原文。', SONG)
save(doc, '03_主题字体.docx', {'majorFont': FZ, 'minorFont': 'SimHei'})

# Fixture 4: no title, hidden/deleted matching font, mixed script fonts, table.
doc = make_doc()
p = para(doc, '隐藏的小标宋不应被选为标题', FZ)
p.runs[0].font.hidden = True
p = doc.add_paragraph()
deleted = OxmlElement('w:del')
deleted.set(qn('w:id'), '1')
deleted.set(qn('w:author'), 'Demo')
r = OxmlElement('w:r')
rp = OxmlElement('w:rPr')
rf = OxmlElement('w:rFonts')
rf.set(qn('w:eastAsia'), FZ)
rp.append(rf)
r.append(rp)
t = OxmlElement('w:delText')
t.text = '删除的小标宋不应被选为标题'
r.append(t)
deleted.append(r)
p._p.append(deleted)
para(doc, '这份文档没有指定的小标宋标题，应使用文件名。', SONG)
para(doc, '一、有效的黑体标题', HEI)
para(doc, '加粗宋体仍然是正文。', SONG, bold=True)
p = doc.add_paragraph()
r = p.add_run('2026 年计划')
font(r, 'Calibri', east=HEI)
para(doc, '上段数字为 Calibri，汉字为黑体，不满足全段黑体。', SONG)
table = doc.add_table(rows=2, cols=2)
table.style = 'Table Grid'
for row, values in zip(table.rows, [('项目', '状态'), ('文档整理', '已完成')]):
    for cell, text in zip(row.cells, values):
        cell.text = text
        for p in cell.paragraphs:
            for r in p.runs:
                font(r, SONG)
save(doc, '04_缺失标题与混合字体.docx')

embedded = [{'name': path.name, 'data': base64.b64encode(path.read_bytes()).decode()} for path in sorted(OUT.glob('*.docx'))]
(ROOT / 'src' / 'demos.json').write_text(json.dumps(embedded, ensure_ascii=False), encoding='utf-8')
with ZipFile(OUT / 'Word转Markdown_示例文档.zip', 'w', ZIP_DEFLATED) as archive:
    for path in sorted(OUT.glob('*.docx')):
        info = ZipInfo(path.name, date_time=(2026, 10, 4, 0, 0, 0))
        info.compress_type = ZIP_DEFLATED
        archive.writestr(info, path.read_bytes())
print('Generated', len(embedded), 'DOCX files and embedded browser demos.')
